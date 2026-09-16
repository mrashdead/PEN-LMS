"""Excel export adapter for the reports read model."""
from __future__ import annotations

from datetime import datetime

from django.http import HttpResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.reports.permissions import CanAccessReports
from apps.reports.selectors.reports import (
    ReportFilterError,
    ReportFilters,
    dashboard_data,
    export_rows,
)


class ReportsExportView(APIView):
    """GET /api/reports/export/?report=all — RTL .xlsx workbook."""

    permission_classes = (IsAuthenticated, CanAccessReports)

    def get(self, request):
        try:
            filters = ReportFilters.from_query(request.query_params)
            report = (request.query_params.get("report") or "all").strip().lower()
            allowed = {"all", "financial", "enrollments", "people", "classes", "workflow", "communications"}
            if report not in allowed:
                return Response({"detail": "نوع گزارش برای خروجی معتبر نیست."}, status=400)
            data = dashboard_data(request.user, filters)
            sheets = export_rows(data, report)
        except ReportFilterError as exc:
            return Response({"detail": str(exc)}, status=400)

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response(
                {"detail": "کتابخانهٔ openpyxl نصب نشده است. requirements.txt را نصب کنید."},
                status=503,
            )

        workbook = Workbook()
        workbook.remove(workbook.active)
        header_fill = PatternFill("solid", fgColor="123B5D")
        header_font = Font(name="Vazirmatn", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="Vazirmatn", size=10, color="17212B")
        thin = Side(style="thin", color="D9E2EC")
        border = Border(bottom=thin)
        right = Alignment(horizontal="right", vertical="center", readingOrder=2)
        center = Alignment(horizontal="center", vertical="center", readingOrder=2)

        for title, headers, rows in sheets:
            sheet = workbook.create_sheet(title=title[:31])
            sheet.sheet_view.rightToLeft = True
            sheet.freeze_panes = "A2"
            sheet.row_dimensions[1].height = 26
            for col, value in enumerate(headers, start=1):
                cell = sheet.cell(row=1, column=col, value=value)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = right
            for row_index, row in enumerate(rows, start=2):
                for col, value in enumerate(row, start=1):
                    cell = sheet.cell(row=row_index, column=col, value=value)
                    cell.font = body_font
                    cell.border = border
                    cell.alignment = center if isinstance(value, (int, float)) else right
            for col in range(1, len(headers) + 1):
                longest = max(
                    [len(str(sheet.cell(row=row, column=col).value or "")) for row in range(1, sheet.max_row + 1)]
                    or [10]
                )
                sheet.column_dimensions[get_column_letter(col)].width = min(max(longest + 4, 14), 38)

        if not sheets:
            sheet = workbook.create_sheet(title="گزارش")
            sheet.sheet_view.rightToLeft = True
            sheet["A1"] = "داده‌ای برای خروجی در دسترس نیست."
            sheet["A1"].font = body_font

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        stamp = datetime.now().strftime("%Y%m%d-%H%M")
        response["Content-Disposition"] = f'attachment; filename="pen-reports-{stamp}.xlsx"'
        workbook.save(response)
        return response
