from io import BytesIO

from django.test import TestCase

from apps.forms.models import FormSchema, FormSubmission, Request, RequestType
from apps.forms.tests.factories import UserFactory
from apps.tasks.models import WorkflowTask
from apps.workflow.models import ActionLog, State
from apps.workflow.tests.factories import InstanceFactory, WorkflowDefinitionFactory


class WorkflowReportingTests(TestCase):
    def test_report_filters_and_exposes_current_assignment_and_form_stage(self):
        manager = UserFactory(username="workflow-report-manager", roles=["manager"])
        requester = UserFactory(username="workflow-report-requester", roles=["employee"])
        assignee = UserFactory(username="workflow-report-assignee", roles=["supervisor"])
        requester.department = "Operations"
        requester.save(update_fields=["department", "updated_at"])
        assignee.department = "Review Unit"
        assignee.save(update_fields=["department", "updated_at"])
        definition = WorkflowDefinitionFactory(code="report-leave", name="Leave")
        state = State.objects.create(
            workflow_definition=definition,
            code="pending-manager",
            name="بررسی مدیر",
        )
        instance = InstanceFactory(
            workflow_definition=definition,
            requester=requester,
            current_state=state,
            title="مرخصی مهر",
        )
        request_type = RequestType.objects.create(
            code="annual-leave-report", title="مرخصی سالانه", workflow_definition=definition,
        )
        schema = FormSchema.objects.create(
            slug="annual-leave-report", title="فرم مرخصی", fields=[],
            request_type=request_type, workflow_definition=definition,
        )
        submission = FormSubmission.objects.create(
            form_schema=schema, submitted_by=requester, status=FormSubmission.Status.SUBMITTED,
            submission_number="FORM-LEAVE-2026-0001", workflow_instance=instance,
        )
        Request.objects.create(
            request_type=request_type, form_submission=submission, requester=requester,
            request_number="REQ-2026-0001", status=Request.Status.IN_REVIEW,
        )
        WorkflowTask.objects.create(
            instance=instance,
            state=state,
            assignee=assignee,
            assigned_by=requester,
            status=WorkflowTask.Status.PENDING,
        )
        ActionLog.objects.create(
            instance=instance,
            action="submit",
            actor=requester,
        )
        ActionLog.objects.create(
            instance=instance,
            action="approve",
            actor=assignee,
            from_state=state,
            to_state=state,
            comment="بررسی شد",
        )
        self.client.force_login(manager)

        response = self.client.get(
            "/api/workflow/reports/?workflow=report-leave&request_type=annual-leave-report"
            "&request_status=in_review&request_department=Review&user=workflow-report-requester"
            "&form=annual-leave-report&from=2026-09-01&to=2026-09-30"
            "&review_from=2026-09-01&review_to=2026-09-30"
        )

        self.assertEqual(response.status_code, 200, response.content)
        row = response.json()["results"][0]
        self.assertEqual(row["id"], str(instance.pk))
        self.assertEqual(row["request_number"], "REQ-2026-0001")
        self.assertEqual(row["request_type_title"], "مرخصی سالانه")
        self.assertEqual(row["form_title"], "فرم مرخصی")
        self.assertEqual(row["current_step"], "بررسی مدیر")
        self.assertEqual(row["assigned_to"], assignee.username)
        self.assertEqual(row["assigned_department"], "Review Unit")
        self.assertEqual(row["action_count"], 2)
        self.assertIsNotNone(row["last_reviewed_at"])
        self.assertIsNotNone(row["approved_at"])
        self.assertEqual(len(row["action_history"]), 2)

        export = self.client.get(
            "/api/reports/export/?report=workflow&request_type=annual-leave-report"
            "&request_department=Review"
        )
        self.assertEqual(export.status_code, 200, export.content[:200])
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export.content), read_only=True)
        self.assertIn("جزئیات درخواست", workbook.sheetnames)
        detail_row = next(workbook["جزئیات درخواست"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(detail_row[0], "REQ-2026-0001")
        self.assertIn("approve", detail_row[-1])

    def test_employee_cannot_read_global_workflow_report(self):
        employee = UserFactory(username="workflow-report-employee", roles=["employee"])
        self.client.force_login(employee)

        response = self.client.get("/api/workflow/reports/")

        self.assertEqual(response.status_code, 403)
